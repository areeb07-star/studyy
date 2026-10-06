#!/usr/bin/env python3
"""Extract page/slide text and render source-linked images for StudyGraph."""

import json
import os
import shutil
import subprocess
import sys
import tempfile

import fitz
from pptx import Presentation

MAX_PAGES = 500
IMAGE_SCALE = 1.35


def clean_text(value: str) -> str:
    return "\n".join(line.strip() for line in value.splitlines() if line.strip())


def pdf_pages(pdf_path: str, output_dir: str, page_type: str, slide_text=None):
    document = fitz.open(pdf_path)
    if document.needs_pass:
        raise RuntimeError("This PDF is password-protected and cannot be processed.")
    if len(document) > MAX_PAGES:
        raise RuntimeError(f"This file exceeds the {MAX_PAGES}-page ingestion limit.")

    pages = []
    for index, page in enumerate(document):
        number = index + 1
        image_name = f"{page_type}-{number:04d}.png"
        image_path = os.path.join(output_dir, image_name)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(IMAGE_SCALE, IMAGE_SCALE), alpha=False)
        pixmap.save(image_path)
        text = slide_text[index] if slide_text and index < len(slide_text) else page.get_text("text")
        pages.append(
            {
                "number": number,
                "text": clean_text(text or ""),
                "image": image_path,
                "type": page_type,
            }
        )
    document.close()
    return pages


def extract_pdf(input_path: str, output_dir: str):
    return pdf_pages(input_path, output_dir, "page")


def extract_pptx(input_path: str, output_dir: str):
    presentation = Presentation(input_path)
    if len(presentation.slides) > MAX_PAGES:
        raise RuntimeError(f"This file exceeds the {MAX_PAGES}-slide ingestion limit.")

    slide_text = []
    for slide in presentation.slides:
        lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        try:
            notes = slide.notes_slide.notes_text_frame.text
            if notes and notes.strip():
                lines.append("Speaker notes: " + notes.strip())
        except (AttributeError, KeyError):
            pass
        slide_text.append("\n".join(lines))

    with tempfile.TemporaryDirectory(prefix="studygraph-lo-") as converted_dir:
        result = subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                converted_dir,
                input_path,
            ],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf_path = os.path.join(
            converted_dir, os.path.splitext(os.path.basename(input_path))[0] + ".pdf"
        )
        if result.returncode != 0 or not os.path.exists(pdf_path):
            message = (result.stderr or result.stdout or "LibreOffice could not render this deck.").strip()
            raise RuntimeError(f"PowerPoint rendering failed: {message[:500]}")
        pages = pdf_pages(pdf_path, output_dir, "slide", slide_text)
    if len(pages) != len(slide_text):
        raise RuntimeError(
            f"Rendered {len(pages)} slides but extracted text from {len(slide_text)} slides."
        )
    return pages


def main():
    if len(sys.argv) != 4:
        raise RuntimeError("Usage: extract.py <pdf|pptx> <input-path> <output-directory>")
    source_type, input_path, output_dir = sys.argv[1:]
    os.makedirs(output_dir, exist_ok=True)
    if source_type == "pdf":
        pages = extract_pdf(input_path, output_dir)
    elif source_type == "pptx":
        pages = extract_pptx(input_path, output_dir)
    else:
        raise RuntimeError(f"Unsupported document type: {source_type}")

    manifest_path = os.path.join(output_dir, "manifest.json")
    with open(manifest_path, "w", encoding="utf-8") as manifest_file:
        json.dump({"pages": pages}, manifest_file, ensure_ascii=False)
    print(json.dumps({"manifestPath": manifest_path, "pageCount": len(pages)}))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
